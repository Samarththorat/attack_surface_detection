import unittest
import os
import tempfile
from io import BytesIO
from unittest.mock import patch

from openpyxl import load_workbook

import database.database as db
from scanner.analytics import build_security_analytics
from scanner.reporting import _safe_cell, build_excel_report


class TestAnalyticsAndReporting(unittest.TestCase):
    def test_vulnerability_metrics_distinguish_observations_and_current_state(self):
        fd, path = tempfile.mkstemp(suffix=".db")
        os.close(fd)
        with patch.object(db, "DB_PATH", path):
            try:
                db.init_db()
                initial = db.get_vulnerability_analytics()
                self.assertEqual(initial["current_assessment_state"], "NO_ASSESSMENT")
                self.assertIsNone(initial["currently_active_findings"])

                confirmed = [{
                    "cve_id": "CVE-2025-1234",
                    "software_name": "Example Product",
                    "version": "1.2.3",
                    "assessment_status": "Confirmed Version Match",
                    "evidence": "Installed version 1.2.3 matches the NVD affected version criteria.",
                    "detected_time": "2025-01-01 12:00:00",
                }]
                potential = [{
                    "cve_id": "CVE-2025-1234",
                    "software_name": "Example Product",
                    "version": "1.2.4",
                    "assessment_status": "Potential Match",
                    "detected_time": "2025-01-02 12:00:00",
                }]
                db.record_vulnerability_assessment_run(1, confirmed, True)
                db.record_vulnerability_assessment_run(1, potential, True)

                analytics = db.get_vulnerability_analytics()
                self.assertEqual(analytics["unique_cves"], 1)
                self.assertEqual(analytics["unique_affected_software_cve_pairs"], 1)
                self.assertEqual(analytics["stored_observations"], 2)
                self.assertEqual(analytics["repeated_observations"], 1)
                self.assertEqual(analytics["confirmed_software_cve_pairs"], 1)
                self.assertEqual(analytics["potential_software_cve_pairs"], 1)
                self.assertEqual(analytics["currently_active_findings"], 1)
                self.assertEqual(analytics["currently_active_potential_pairs"], 1)

                db.record_vulnerability_assessment_run(1, [], False)
                incomplete = db.get_vulnerability_analytics()
                self.assertEqual(incomplete["current_assessment_state"], "INCOMPLETE")
                self.assertIsNone(incomplete["currently_active_findings"])
            finally:
                os.remove(path)

    def test_success_rate_excludes_unverified_actions(self):
        remediation = [
            {"verification_status": "RESOLVED", "status": "SUCCESS"},
            {"verification_status": "PARTIALLY_RESOLVED", "status": "SUCCESS"},
            {"verification_status": "FAILED", "status": "FAILED"},
            {"verification_status": "UNVERIFIED", "status": "SUCCESS"},
            {"status": "SUCCESS"},
        ]
        with (
            patch("scanner.analytics.get_scan_history", return_value=[]),
            patch("scanner.analytics.get_remediation_history", return_value=remediation),
            patch("scanner.analytics.get_finding_lifecycles", return_value=[]),
            patch("scanner.analytics.get_monitoring_events", return_value=[]),
            patch(
                "scanner.analytics.get_vulnerability_analytics",
                return_value={
                    "vulnerability_findings_by_day": {},
                    "unique_cves": 0,
                    "unique_affected_software_cve_pairs": 0,
                    "stored_observations": 0,
                    "repeated_observations": 0,
                    "confirmed_software_cve_pairs": 0,
                    "potential_software_cve_pairs": 0,
                    "current_assessment_state": "NO_ASSESSMENT",
                    "currently_active_findings": None,
                },
            ),
        ):
            analytics = build_security_analytics()

        self.assertEqual(analytics["remediation_success_rate"], 33.33)

    def test_report_cell_escaping_handles_whitespace_prefixed_formulas(self):
        malicious = " \t=HYPERLINK(\"https://example.invalid\")"

        self.assertEqual(_safe_cell(malicious), "'" + malicious)
        self.assertEqual(_safe_cell("-1"), "'-1")
        self.assertEqual(_safe_cell("ordinary text"), "ordinary text")

    def test_excel_report_is_a_readable_workbook_with_empty_sections(self):
        with (
            patch("scanner.reporting.get_scan_history", return_value=[]),
            patch("scanner.reporting.get_scan_by_id", return_value=None),
            patch("scanner.reporting.get_remediation_history", return_value=[]),
            patch("scanner.reporting.get_monitoring_events", return_value=[]),
            patch("scanner.reporting.get_vulnerability_findings", return_value=[]),
            patch("scanner.reporting.get_system_monitor_samples", return_value=[]),
            patch(
                "scanner.reporting.build_security_analytics",
                return_value={
                    "risk_trend": [],
                    "vulnerability_findings_by_day": {},
                },
            ),
        ):
            report = build_excel_report()

        workbook = load_workbook(BytesIO(report.getvalue()), read_only=True)
        self.assertEqual(
            workbook.sheetnames,
            [
                "Executive Summary",
                "Scan History",
                "Security Findings",
                "Risk Analysis",
                "Vulnerability Findings",
                "Monitoring Events",
                "Remediation History",
                "Verification Results",
                "Process and System Monitoring",
                "Security Trends",
                "Methodology and Limitations",
            ],
        )

    def test_excel_report_marks_partial_latest_scan(self):
        scan = {
            "id": 7,
            "scan_time": "2025-01-01 12:00:00",
            "risk_score": 20,
            "risk_level": "Low",
            "full_results": {
                "assessment_status": "PARTIAL",
                "incomplete_sections": ["firewall"],
                "risk": {"unified_findings": []},
            },
        }
        with (
            patch("scanner.reporting.get_scan_history", return_value=[{"id": 7}]),
            patch("scanner.reporting.get_scan_by_id", return_value=scan),
            patch("scanner.reporting.get_remediation_history", return_value=[]),
            patch("scanner.reporting.get_monitoring_events", return_value=[]),
            patch("scanner.reporting.get_vulnerability_findings", return_value=[]),
            patch("scanner.reporting.get_system_monitor_samples", return_value=[]),
            patch(
                "scanner.reporting.build_security_analytics",
                return_value={
                    "risk_trend": [],
                    "vulnerability_findings_by_day": {},
                    "vulnerability_summary": {},
                    "methodology": {},
                },
            ),
        ):
            report = build_excel_report()

        workbook = load_workbook(BytesIO(report.getvalue()), read_only=True)
        summary = list(workbook["Executive Summary"].values)
        summary_by_name = dict(zip(summary[0], summary[1]))
        self.assertEqual(summary_by_name["latest_scan_assessment_status"], "PARTIAL")
        self.assertEqual(summary_by_name["report_completeness"], "PARTIAL")

    def test_excel_report_exports_actual_system_samples_and_linked_verification_scan(self):
        sample = {
            "sampled_at": "2025-01-02T03:04:05+00:00",
            "cpu_usage_percent": 12.5,
            "running_processes": 42,
        }
        remediation = {
            "id": 8,
            "action_type": "BLOCK_PORT",
            "scan_id": 3,
            "verification_scan_id": 4,
            "verification_status": "RESOLVED",
        }
        with (
            patch("scanner.reporting.get_scan_history", return_value=[]),
            patch("scanner.reporting.get_scan_by_id", return_value=None),
            patch("scanner.reporting.get_remediation_history", return_value=[remediation]),
            patch("scanner.reporting.get_monitoring_events", return_value=[]),
            patch("scanner.reporting.get_vulnerability_findings", return_value=[]),
            patch("scanner.reporting.get_system_monitor_samples", return_value=[sample]),
            patch(
                "scanner.reporting.build_security_analytics",
                return_value={"risk_trend": [], "vulnerability_findings_by_day": {}},
            ),
        ):
            report = build_excel_report()

        workbook = load_workbook(BytesIO(report.getvalue()), read_only=True, data_only=True)
        summary = list(workbook["Executive Summary"].values)
        summary_by_name = dict(zip(summary[0], summary[1]))
        self.assertEqual(summary_by_name["system_monitor_samples_exported"], 1)

        sample_sheet = list(workbook["Process and System Monitoring"].values)
        self.assertEqual(sample_sheet[1][sample_sheet[0].index("cpu_usage_percent")], 12.5)
        remediation_sheet = list(workbook["Remediation History"].values)
        self.assertEqual(remediation_sheet[1][remediation_sheet[0].index("verification_scan_id")], 4)
        workbook.close()

    def test_excel_report_identifies_malformed_stored_scan_sections(self):
        scan = {
            "id": 9,
            "scan_time": "2025-01-01 12:00:00",
            "risk_score": 20,
            "risk_level": "Low",
            "full_results": {
                "assessment_status": "COMPLETE",
                "risk": {
                    "unified_findings": {"unexpected": "object"},
                    "breakdown": {"items": "unexpected string"},
                },
            },
        }
        with (
            patch("scanner.reporting.get_scan_history", return_value=[{"id": 9}]),
            patch("scanner.reporting.get_scan_by_id", return_value=scan),
            patch("scanner.reporting.get_remediation_history", return_value=[]),
            patch("scanner.reporting.get_monitoring_events", return_value=[]),
            patch("scanner.reporting.get_vulnerability_findings", return_value=[]),
            patch("scanner.reporting.get_system_monitor_samples", return_value=[]),
            patch(
                "scanner.reporting.build_security_analytics",
                return_value={
                    "risk_trend": [],
                    "vulnerability_findings_by_day": {},
                    "vulnerability_summary": {},
                    "methodology": {},
                },
            ),
        ):
            report = build_excel_report()

        workbook = load_workbook(BytesIO(report.getvalue()), read_only=True)
        summary = list(workbook["Executive Summary"].values)
        summary_by_name = dict(zip(summary[0], summary[1]))
        self.assertEqual(summary_by_name["report_completeness"], "PARTIAL")
        self.assertIn("Security Findings", summary_by_name["report_completeness_notes"])
        self.assertIn("Risk Analysis", summary_by_name["report_completeness_notes"])


if __name__ == "__main__":
    unittest.main()
