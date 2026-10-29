"""Generate in-memory security reports from stored assessment data."""

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from database.database import (
    get_finding_lifecycles,
    get_monitoring_events,
    get_remediation_history,
    get_scan_by_id,
    get_scan_history,
    get_system_monitor_samples,
    get_vulnerability_findings,
)
from scanner.analytics import build_security_analytics


def _safe_cell(value):
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=True, sort_keys=True)
    if isinstance(value, str):
        formula_probe = value
        while formula_probe and (
            formula_probe[0].isspace()
            or ord(formula_probe[0]) < 32
            or formula_probe[0] == "\ufeff"
        ):
            formula_probe = formula_probe[1:]
        if formula_probe.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def _rows_for_sheet(sheet, rows):
    rows = list(rows or [])
    columns = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    if not columns:
        columns = ["No stored data"]
        sheet.append(columns)
        sheet.append(["No records are available for this report section."])
        return

    sheet.append(columns)
    for row in rows:
        sheet.append([_safe_cell(row.get(column)) for column in columns])


def _validated_rows(value, section, issues):
    if not isinstance(value, (list, tuple)):
        if value not in (None, []):
            issues.append(f"{section} contains malformed data")
        return []
    rows = [row for row in value if isinstance(row, dict)]
    if len(rows) != len(value):
        issues.append(f"{section} contains malformed rows")
    return rows


def _format_sheet(sheet):
    sheet.freeze_panes = "A2"
    if sheet.max_row >= 1 and sheet.max_column >= 1:
        sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in sheet.columns:
        cells = list(column)
        width = min(48, max(12, max(len(str(cell.value or "")) for cell in cells) + 2))
        sheet.column_dimensions[get_column_letter(cells[0].column)].width = width
    sheet.row_dimensions[1].height = 28


def build_excel_report():
    """Return an Excel workbook containing only data already present in SQLite."""
    exported_sections = []
    history = _validated_rows(get_scan_history(limit=500), "Scan History", exported_sections)
    latest_scan = None
    if history:
        latest_id = history[0].get("id")
        if latest_id is None:
            exported_sections.append("Latest scan history entry has no ID")
        else:
            latest_scan = get_scan_by_id(latest_id)
    if latest_scan is not None and not isinstance(latest_scan, dict):
        exported_sections.append("Latest scan record contains malformed data")
        latest_scan = None
    latest_results = (latest_scan or {}).get("full_results") or {}
    if not isinstance(latest_results, dict):
        exported_sections.append("Latest scan payload contains malformed data")
        latest_results = {}
    raw_risk = latest_results.get("risk")
    risk = raw_risk if isinstance(raw_risk, dict) else {}
    if raw_risk is not None and not isinstance(raw_risk, dict):
        exported_sections.append("Latest scan risk data contains malformed data")
    analytics = build_security_analytics()
    if not isinstance(analytics, dict):
        exported_sections.append("Analytics contains malformed data")
        analytics = {}
    remediation = _validated_rows(
        get_remediation_history(limit=1000),
        "Remediation History",
        exported_sections,
    )
    system_monitoring = _validated_rows(
        get_system_monitor_samples(hours=168, limit=10080),
        "Process and System Monitoring",
        exported_sections,
    )
    vulnerability_rows = _validated_rows(
        get_vulnerability_findings(limit=1000),
        "Vulnerability Findings",
        exported_sections,
    )
    for row in vulnerability_rows:
        recorded_status = row.get("assessment_status")
        row["recorded_assessment_status"] = recorded_status
        row["assessment_status"] = row.get("effective_assessment_status") or recorded_status
    vulnerability_summary = analytics.get("vulnerability_summary", {})
    if not isinstance(vulnerability_summary, dict):
        exported_sections.append("Vulnerability summary contains malformed data")
        vulnerability_summary = {}
    latest_scan_status = latest_results.get("assessment_status")
    if latest_scan_status not in ("COMPLETE", "PARTIAL"):
        latest_scan_status = "UNKNOWN (legacy scan has no completeness metadata)" if latest_scan else "NO SCAN"
    raw_findings = risk.get("unified_findings", [])
    findings = _validated_rows(raw_findings, "Security Findings", exported_sections)

    raw_breakdown = risk.get("breakdown")
    raw_risk_rows = raw_breakdown.get("items", []) if isinstance(raw_breakdown, dict) else []
    risk_rows = _validated_rows(raw_risk_rows, "Risk Analysis", exported_sections)
    if raw_breakdown is not None and not isinstance(raw_breakdown, dict):
        exported_sections.append("Risk Analysis breakdown contains malformed data")

    risk_trend = _validated_rows(analytics.get("risk_trend", []), "Risk Trend", exported_sections)
    vulnerability_trend = analytics.get("vulnerability_findings_by_day", {})
    if not isinstance(vulnerability_trend, dict):
        exported_sections.append("Vulnerability trend contains malformed data")
        vulnerability_trend = {}
    methodology = analytics.get("methodology", {})
    if not isinstance(methodology, dict):
        exported_sections.append("Analytics methodology contains malformed data")
        methodology = {}

    if latest_scan_status != "COMPLETE":
        exported_sections.append(f"Latest scan: {latest_scan_status}")
    if len(history) >= 500:
        exported_sections.append("Scan History is limited to the latest 500 records")
    if len(vulnerability_rows) >= 1000:
        exported_sections.append("Vulnerability Findings is limited to the latest 1000 records")
    if not system_monitoring:
        exported_sections.append("No stored process/system monitoring samples are available")
    elif len(system_monitoring) >= 10080:
        exported_sections.append("Process and System Monitoring is limited to the latest 10080 samples")
    report_status = "PARTIAL" if exported_sections else "COMPLETE"

    summary_rows = [{
        "generated_at": __import__("datetime").datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "latest_scan_id": (latest_scan or {}).get("id"),
        "latest_scan_time": (latest_scan or {}).get("scan_time"),
        "latest_risk_score": risk.get("total_score", (latest_scan or {}).get("risk_score")),
        "latest_risk_level": risk.get("risk_level", (latest_scan or {}).get("risk_level")),
        "latest_findings_count": len(findings),
        "latest_scan_assessment_status": latest_scan_status,
        "latest_scan_incomplete_sections": latest_results.get("incomplete_sections", []),
        "scan_records_included": len(history),
        "report_completeness": report_status,
        "report_completeness_notes": exported_sections,
        "vulnerability_records_exported": len(vulnerability_rows),
        "system_monitor_samples_exported": len(system_monitoring),
        "vulnerability_observations_stored": vulnerability_summary.get("stored_observations", 0),
        "latest_vulnerability_assessment_state": vulnerability_summary.get("current_assessment_state", "NO_ASSESSMENT"),
        "currently_active_vulnerability_findings": vulnerability_summary.get("currently_active_findings"),
    }]
    trends = [
        {"type": "risk_score", **row}
        for row in risk_trend
    ] + [
        {"type": "vulnerability_findings_by_day", "date": date, "count": count}
        for date, count in vulnerability_trend.items()
    ]
    methodology_rows = [
        {"topic": "Risk score", "method": "Stored project-specific 0-100 heuristic. See score breakdown; not a certified rating."},
        {"topic": "Finding resolution", "method": "Only evidence from comparable complete stored scans is used for absence-based comparisons; missing/incomplete sections are not resolved."},
        {"topic": "Vulnerability assessment", "method": "NVD records matched to installed software identity and available affected-version data. A version match does not prove exploitability."},
        {"topic": "Assessment completeness", "method": f"Latest scan status: {latest_scan_status}; incomplete sections: {latest_results.get('incomplete_sections', []) or 'none recorded'}. Report status: {report_status}; notes: {exported_sections or 'none'}."},
        {"topic": "Vulnerability counts", "method": methodology.get("vulnerability_summary", "Unique and active counts are not available.")},
        {"topic": "Internet exposure", "method": "A local listening socket does not prove internet accessibility."},
        {"topic": "Process and system monitoring", "method": "This sheet contains stored system-level samples only. GPU utilization, per-process network attribution, and per-process energy are unavailable and are not estimated. Process-level snapshots are live only and are not retained."},
        {"topic": "Report limitations", "method": "Only records already stored in this local database are included. This report does not synthesize missing history."},
        {"topic": "Sensitive data", "method": "API credentials and tokens are not exported."},
    ]

    sheet_data = [
        ("Executive Summary", summary_rows),
        ("Scan History", history),
        ("Security Findings", findings),
        ("Risk Analysis", risk_rows),
        ("Vulnerability Findings", vulnerability_rows),
        ("Monitoring Events", _validated_rows(
            get_monitoring_events(limit=1000),
            "Monitoring Events",
            exported_sections,
        )),
        ("Remediation History", remediation),
        ("Verification Results", [
            row for row in remediation
            if row.get("verification_status") or row.get("actual_outcome")
        ]),
        ("Process and System Monitoring", system_monitoring),
        ("Security Trends", trends),
        ("Methodology and Limitations", methodology_rows),
    ]

    workbook = Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheet_data:
        sheet = workbook.create_sheet(title=title)
        _rows_for_sheet(sheet, rows)
        _format_sheet(sheet)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
